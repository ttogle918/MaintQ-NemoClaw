# -*- coding: utf-8 -*-
"""docx_render — 결재 문서 템플릿을 채워 **메모리에서만** .docx bytes 를 만든다 (D124).

★ 저장하지 않는다 (D86)
────────────────────────────────────────────────────────────────────────────────
원본 템플릿을 읽어 `io.BytesIO` 에 쓰고 그 bytes 를 돌려준다. 디스크에도 DB 에도 아무것도
남기지 않는다 — 임시 파일조차 만들지 않는다. 저장하면 템플릿·산식이 바뀔 때 저장본이
조용히 낡는다(D57·D86 선례).

★ 양방향 strict — 왜 관대하게 만들지 않는가
────────────────────────────────────────────────────────────────────────────────
  · 템플릿에 있는데 fields 에 없는 키 → 그대로 두면 **결재 서류에 `{{X}}` 가 인쇄된다.**
    테스트가 아니라 코드가 막아야 하는 종류다.
  · fields 에 있는데 템플릿에 없는 키 → 템플릿과 코드가 갈렸다는 뜻이다. 조용히 무시하면
    아무도 모른다 — D90 이 겪은 "같은 문장이 두 곳에 있어 한쪽만 바뀐다"와 같은 유형이다.

  ⚠ 단 **`drop_rows` 로 사라진 자리는 초과로 보지 않는다.** 호출부가 조건부로 행을 지울 때
    fields 쪽도 함께 좁히도록 강요하면 두 곳을 같이 고쳐야 하는 구조가 되는데, 그게 바로
    위에서 막으려는 실패 유형이다.

★ 런 분할을 조용히 처리하지 않는다
────────────────────────────────────────────────────────────────────────────────
2026-09-03 실측에서 5종 템플릿 전부 플레이스홀더가 단일 `<w:t>` 안에 온전했다(01:46 ·
02:55 · 03:41 · 05:30 · 06:27, 분할 0건). 그래서 run-merge 로직을 선제적으로 넣지 않는다.
대신 분할을 **발견하면 명시적으로 실패**시킨다 — 조용히 넘기면 나중에 템플릿이 편집돼
분할이 생겼을 때 값이 안 채워진 문서가 결재에 올라간다.

★ 대상은 5종이다
────────────────────────────────────────────────────────────────────────────────
01·02·03·05·06. **04(담보대출심사회신서)는 만들지 않는다** — D118 이 MaintQ 구현 대상에서
제외했다(FinAllQ 소관). 파일은 `data/templates/` 에 있지만 채우는 코드는 없다.
"""

from __future__ import annotations

import io
import re
from collections.abc import Iterable, Iterator, Mapping
from pathlib import Path

from docx import Document

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

TEMPLATE_DIR = Path(__file__).resolve().parents[2] / "data" / "templates"

_PLACEHOLDER = re.compile(r"\{\{([A-Z0-9_]+)\}\}")


class TemplateFieldMismatch(Exception):
    """템플릿 자리와 fields 키가 어긋났다. 어느 쪽이 남았는지 메시지에 적는다."""


class SplitPlaceholderError(Exception):
    """플레이스홀더가 여러 run 으로 쪼개져 있다 — 치환하면 값이 안 들어간다."""


def _template_path(template_filename: str) -> Path:
    path = TEMPLATE_DIR / template_filename
    if not path.is_file():
        raise FileNotFoundError(f"템플릿이 없습니다: {path}")
    return path


def _iter_tables(container) -> Iterator:
    for table in container.tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_tables(cell)


def _iter_paragraphs(container) -> Iterator:
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _all_paragraphs(doc) -> Iterator:
    """본문 + 표(중첩 포함) + 머리글·바닥글. 자리가 어디 있든 놓치지 않는다."""
    yield from _iter_paragraphs(doc)
    for section in doc.sections:
        yield from _iter_paragraphs(section.header)
        yield from _iter_paragraphs(section.footer)


def _all_tables(doc) -> Iterator:
    yield from _iter_tables(doc)
    for section in doc.sections:
        yield from _iter_tables(section.header)
        yield from _iter_tables(section.footer)


def _found(doc) -> set[str]:
    keys: set[str] = set()
    for p in _all_paragraphs(doc):
        keys |= set(_PLACEHOLDER.findall(p.text))
    return keys


def _assert_not_split(doc) -> None:
    for p in _all_paragraphs(doc):
        joined = set(_PLACEHOLDER.findall(p.text))
        if not joined:
            continue
        per_run: set[str] = set()
        for run in p.runs:
            per_run |= set(_PLACEHOLDER.findall(run.text))
        if split := joined - per_run:
            raise SplitPlaceholderError(
                f"플레이스홀더가 여러 run 으로 쪼개져 있습니다: {sorted(split)}"
            )


def template_placeholders(template_filename: str) -> set[str]:
    """템플릿이 요구하는 자리 전부.

    계약 검증과 테스트가 기대값을 **여기서 파생시킨다** — 목록을 하드코딩하면 템플릿이
    바뀔 때 기대값이 같이 안 움직여 회귀가 거짓 신호를 낸다.
    """
    return _found(Document(str(_template_path(template_filename))))


def _drop_rows(doc, keys: set[str]) -> set[str]:
    """`keys` 중 하나라도 든 표 행을 삭제한다. 삭제로 사라진 자리 집합을 돌려준다.

    빈 칸을 `확인되지 않음` 으로 채우지 않고 행 자체를 없애는 이유는 미리보기와 같다 —
    `po_documents.py` 가 *"없는 2·3행을 빈 칸으로 채우지 않고 아예 생략한다"* 고 못박았다.
    서류의 빈 칸은 읽는 사람에게 **'해당 없음'으로 읽힌다**(D62).
    """
    removed: set[str] = set()
    for table in _all_tables(doc):
        for row in list(table.rows):
            in_row = set(_PLACEHOLDER.findall("".join(c.text for c in row.cells)))
            if in_row & keys:
                removed |= in_row
                row._element.getparent().remove(row._element)
    return removed


def fill_template(
    template_filename: str,
    fields: Mapping[str, str],
    *,
    drop_rows: Iterable[str] = (),
) -> bytes:
    """템플릿을 채워 .docx bytes 를 돌려준다. 어디에도 저장하지 않는다 (D86).

    `drop_rows` 에 든 자리를 품은 표 행은 채우기 전에 삭제된다 — 그 행의 자리들은
    `fields` 에 없어도 되고, 있어도 초과로 보지 않는다.
    """
    doc = Document(str(_template_path(template_filename)))
    _assert_not_split(doc)

    dropped = _drop_rows(doc, set(drop_rows)) if drop_rows else set()
    required = _found(doc)
    given = set(fields)

    if missing := required - given:
        raise TemplateFieldMismatch(
            f"{template_filename}: 템플릿에 있는데 값이 없는 자리 {sorted(missing)}"
        )
    if extra := given - required - dropped:
        raise TemplateFieldMismatch(
            f"{template_filename}: 값은 있는데 템플릿에 없는 키 {sorted(extra)}"
        )

    for p in _all_paragraphs(doc):
        for run in p.runs:
            if "{{" in run.text:
                run.text = _PLACEHOLDER.sub(lambda m: fields[m.group(1)], run.text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
